"""Étape 7 : recensement.

Produit le classeur Excel de contrôle remis à la cliente avec la base :
résumé chiffré, inventaire complet, doublons, documents retirés, et liste
« à vérifier ». Chaque ligne dit quelle règle a décidé : tout est relisible.
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import KEEP, Attachment, Category

COLUMNS = [
    "final_name", "category", "entity", "year", "confidence", "reason",
    "mailbox", "mail_date", "direction", "sender_domain", "has_email", "has_phone",
    "duplicate_of", "sha256",
]
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")


def _sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    for r in rows:
        ws.append([str(r.get(k, "")) for k in COLUMNS])
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = 22 if col != "reason" else 38
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def summary(kept: list[Attachment], dups: list[Attachment], n_mails: int = 0) -> dict:
    cats = Counter(a.category.value for a in kept)
    in_base = sum(n for c, n in cats.items() if Category(c) in KEEP)
    return {
        "mails_lus": n_mails,
        "pieces_jointes_extraites": len(kept) + len(dups),
        "doublons_ecartes": len(dups),
        "documents_dans_la_base": in_base,
        "documents_retires": len(kept) - in_base - cats.get("a_verifier", 0),
        "a_verifier": cats.get("a_verifier", 0),
        **{f"cat_{k}": v for k, v in sorted(cats.items())},
    }


def write_report(path: Path, kept: list[Attachment], dups: list[Attachment], n_mails: int = 0) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "Résumé"
    stats = summary(kept, dups, n_mails)
    ws.append(["Indicateur", "Valeur"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    for k, v in stats.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 32

    in_base = [a.to_row() for a in kept if a.category in KEEP]
    removed = [a.to_row() for a in kept if a.category not in KEEP | {Category.A_VERIFIER}]
    check = [a.to_row() for a in kept if a.category is Category.A_VERIFIER]
    _sheet(wb, "Base", in_base)
    _sheet(wb, "À vérifier", check)
    _sheet(wb, "Retirés", removed)
    _sheet(wb, "Doublons", [a.to_row() for a in dups])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return stats
