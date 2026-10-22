"""Test de bout en bout sur des boîtes mail fictives."""

from openpyxl import load_workbook

from mailrecovery.pipeline import Source, run
from mailrecovery.synthetic import generate


def test_end_to_end(tmp_path):
    boxes = generate(tmp_path / "src", n_mails=120, seed=7)
    out = tmp_path / "BASE"
    stats = run([Source(b.name, b) for b in boxes], out, own_domains={"cabinet-demo.fr"}, convert=False)

    assert stats["mails_lus"] == 120
    assert stats["documents_dans_la_base"] > 0
    assert stats["doublons_ecartes"] > 0

    cvs = list((out / "1_BASE_CANDIDATS" / "CV_complets").iterdir())
    assert cvs and all(p.stem.split()[0].isupper() for p in cvs)    # « NOM Prénom »
    assert (out / "_RETIRES").exists()

    wb = load_workbook(out / "RECENSEMENT.xlsx")
    assert {"Résumé", "Base", "À vérifier", "Retirés", "Doublons"} <= set(wb.sheetnames)
    # rien n'est perdu : chaque pièce jointe est soit dans un dossier, soit en doublon
    n_files = sum(1 for p in out.rglob("*") if p.is_file() and p.name != "RECENSEMENT.xlsx")
    assert n_files == stats["pieces_jointes_extraites"]
